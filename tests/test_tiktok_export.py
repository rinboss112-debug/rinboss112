from __future__ import annotations

import io
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from streamlit.testing.v1 import AppTest

from tiktok_export import (
    build_preparation_workbook,
    calculate_tiktok_price,
    detect_header_row,
    fill_official_template,
    prepare_products,
    suggest_template_mapping,
    template_headers,
    validate_products,
    workbook_sheet_names,
)


def _sample_source() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "keyword": "snack box",
                "asin": "B012345678",
                "title": "Healthy Snack Box Variety Pack for Adults",
                "price": 10.0,
                "image_url": "https://example.com/snack.jpg",
                "product_url": "https://www.amazon.com/dp/B012345678",
            }
        ]
    )


def _template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Upload"
    sheet.append(["TikTok Shop US category template"])
    sheet.append(
        [
            "Seller SKU",
            "Product Title",
            "Product Description",
            "Brand",
            "Retail Price",
            "Quantity",
            "Main Image",
        ]
    )
    for column in range(1, 8):
        sheet.cell(3, column).fill = PatternFill("solid", fgColor="FFF7ED")
    notes = workbook.create_sheet("Instructions")
    notes["A1"] = "Do not delete this sheet"
    notes["B2"] = "=1+1"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class _FakeController:
    def snapshot(self) -> dict[str, object]:
        return {"results": _sample_source().to_dict(orient="records")}


class TikTokExportTests(unittest.TestCase):
    def test_price_formula(self) -> None:
        self.assertEqual(calculate_tiktok_price(10), 40.0)
        self.assertEqual(calculate_tiktok_price(15), 52.5)
        self.assertEqual(calculate_tiktok_price(20), 65.0)

    def test_prepare_products_maps_scraper_columns(self) -> None:
        frame = prepare_products(_sample_source())
        self.assertEqual(frame.loc[0, "seller_sku"], "SNACK_BOX-B012345678")
        self.assertEqual(frame.loc[0, "calculated_price"], 40.0)
        self.assertEqual(frame.loc[0, "tiktok_price"], 40.0)
        self.assertIn("Thiếu mô tả", frame.loc[0, "validation"])
        self.assertIn("Chưa xác nhận brand", frame.loc[0, "validation"])

    def test_preparation_workbook_contains_formula_and_settings(self) -> None:
        frame = prepare_products(_sample_source())
        raw = build_preparation_workbook(frame)
        workbook = load_workbook(io.BytesIO(raw), data_only=False)
        try:
            self.assertEqual(
                workbook["TikTok Products"]["I2"].value,
                "=ROUND((H2*'Pricing Settings'!$B$2+'Pricing Settings'!$B$3)/"
                "'Pricing Settings'!$B$4,2)",
            )
            self.assertEqual(workbook["Pricing Settings"]["B2"].value, 2.0)
            self.assertEqual(workbook["Pricing Settings"]["B3"].value, 12.0)
            self.assertEqual(workbook["Pricing Settings"]["B4"].value, 0.8)
        finally:
            workbook.close()

    def test_validation_flags_duplicate_seller_sku(self) -> None:
        source = pd.concat([_sample_source(), _sample_source()], ignore_index=True)
        frame = validate_products(prepare_products(source))
        self.assertIn("Trùng Seller SKU", frame.loc[0, "validation"])
        self.assertIn("Trùng Seller SKU", frame.loc[1, "validation"])

    def test_fill_official_template_preserves_other_sheet(self) -> None:
        template = _template_bytes()
        self.assertEqual(workbook_sheet_names(template), ["Upload", "Instructions"])
        header_row = detect_header_row(template, "Upload")
        self.assertEqual(header_row, 2)
        headers = template_headers(template, "Upload", header_row)
        suggestions = suggest_template_mapping(headers)
        self.assertEqual(suggestions["seller_sku"], "Seller SKU")
        self.assertEqual(suggestions["product_title"], "Product Title")
        self.assertEqual(suggestions["tiktok_price"], "Retail Price")

        products = prepare_products(_sample_source())
        products.loc[0, "description"] = "A varied snack pack for gifting and sharing."
        products.loc[0, "brand"] = "Authorized Test Brand"
        filled = fill_official_template(
            template,
            sheet_name="Upload",
            header_row=header_row,
            mapping=suggestions,
            products=products,
        )
        workbook = load_workbook(io.BytesIO(filled), data_only=False)
        try:
            upload = workbook["Upload"]
            self.assertEqual(upload["A3"].value, "SNACK_BOX-B012345678")
            self.assertEqual(
                upload["B3"].value,
                "Healthy Snack Box Variety Pack for Adults",
            )
            self.assertEqual(upload["E3"].value, 40.0)
            self.assertEqual(upload["A3"].fill.fgColor.rgb, "00FFF7ED")
            self.assertEqual(
                workbook["Instructions"]["A1"].value,
                "Do not delete this sheet",
            )
            self.assertEqual(workbook["Instructions"]["B2"].value, "=1+1")
        finally:
            workbook.close()

    def test_tiktok_page_renders_for_authenticated_admin(self) -> None:
        project_dir = Path(__file__).resolve().parents[1]
        app = AppTest.from_file(
            project_dir / "app_pages" / "tiktok_export.py",
            default_timeout=30,
        )
        app.secrets = {"admin": {"password": "test"}}
        app.session_state["admin_authenticated"] = True
        app.session_state["controller"] = _FakeController()
        app.run()
        self.assertEqual([], list(app.exception))
        self.assertEqual(len(app.metric), 4)


if __name__ == "__main__":
    unittest.main()
