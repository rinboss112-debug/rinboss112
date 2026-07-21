from __future__ import annotations

import io
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from amazon_scraper import slugify_filename


TIKTOK_XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

EDITOR_COLUMNS = [
    "selected",
    "keyword",
    "asin",
    "seller_sku",
    "product_title",
    "description",
    "brand",
    "original_price",
    "calculated_price",
    "tiktok_price",
    "stock",
    "image_1",
    "image_2",
    "product_url",
    "package_weight_lb",
    "package_length_in",
    "package_width_in",
    "package_height_in",
    "gtin_upc",
    "variation_name",
    "variation_value",
    "validation",
]


@dataclass(frozen=True, slots=True)
class ExportField:
    key: str
    label: str
    aliases: tuple[str, ...]


EXPORT_FIELDS = (
    ExportField("seller_sku", "Seller SKU", ("seller sku", "sku", "sku id")),
    ExportField(
        "product_title",
        "Tên sản phẩm",
        ("product title", "product name", "title", "name"),
    ),
    ExportField(
        "description",
        "Mô tả sản phẩm",
        ("product description", "description", "desc"),
    ),
    ExportField("brand", "Thương hiệu", ("brand", "brand name")),
    ExportField(
        "tiktok_price",
        "Giá bán TikTok",
        ("retail price", "sale price", "sales price", "price", "unit price"),
    ),
    ExportField(
        "stock",
        "Tồn kho",
        ("quantity", "stock", "inventory", "available stock"),
    ),
    ExportField(
        "image_1",
        "Ảnh chính",
        ("main image", "main product image", "product image 1", "image 1"),
    ),
    ExportField(
        "image_2",
        "Ảnh phụ 2",
        ("product image 2", "image 2", "second image"),
    ),
    ExportField(
        "package_weight_lb",
        "Khối lượng đóng gói (lb)",
        ("package weight", "weight", "package weight (lb)", "weight (lb)"),
    ),
    ExportField(
        "package_length_in",
        "Chiều dài (in)",
        ("package length", "length", "package length (in)", "length (in)"),
    ),
    ExportField(
        "package_width_in",
        "Chiều rộng (in)",
        ("package width", "width", "package width (in)", "width (in)"),
    ),
    ExportField(
        "package_height_in",
        "Chiều cao (in)",
        ("package height", "height", "package height (in)", "height (in)"),
    ),
    ExportField(
        "gtin_upc",
        "GTIN/UPC",
        ("gtin", "upc", "gtin/upc", "product identifier"),
    ),
    ExportField(
        "variation_name",
        "Tên biến thể 1",
        ("variation 1 name", "variation name", "sales attribute 1"),
    ),
    ExportField(
        "variation_value",
        "Giá trị biến thể 1",
        ("variation 1 value", "variation value", "sales attribute value 1"),
    ),
)


def calculate_tiktok_price(
    original_price: float | int | None,
    multiplier: float = 2.0,
    fixed_cost: float = 12.0,
    divisor: float = 0.8,
) -> float | None:
    if original_price is None or pd.isna(original_price):
        return None
    price = float(original_price)
    if price < 0:
        return None
    if divisor <= 0:
        raise ValueError("Hệ số chia phải lớn hơn 0.")
    return round((price * float(multiplier) + float(fixed_cost)) / divisor, 2)


def read_product_file(filename: str, raw: bytes) -> pd.DataFrame:
    suffix = Path(filename).suffix.casefold()
    if suffix == ".xlsx":
        frame = pd.read_excel(io.BytesIO(raw))
    elif suffix == ".csv":
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
            try:
                frame = pd.read_csv(
                    io.BytesIO(raw), encoding=encoding, sep=None, engine="python"
                )
                break
            except (UnicodeDecodeError, pd.errors.ParserError) as error:
                last_error = error
        else:
            raise ValueError(f"Không đọc được file CSV: {last_error}")
    else:
        raise ValueError("Chỉ hỗ trợ file CSV hoặc XLSX.")
    frame.columns = [str(column).strip() for column in frame.columns]
    return frame.dropna(how="all").reset_index(drop=True)


def _normalized_name(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).strip().casefold()).strip()


def _source_series(
    frame: pd.DataFrame,
    aliases: Iterable[str],
    default: Any = "",
) -> pd.Series:
    by_normalized = {_normalized_name(column): column for column in frame.columns}
    for alias in aliases:
        column = by_normalized.get(_normalized_name(alias))
        if column is not None:
            return frame[column]
    return pd.Series([default] * len(frame), index=frame.index)


def _clean_cell(value: object) -> object:
    if value is None or pd.isna(value):
        return ""
    return value


def _build_sku(keyword: object, asin: object, row_number: int) -> str:
    niche = slugify_filename(str(_clean_cell(keyword)))[:12].upper() or "PRODUCT"
    clean_asin = re.sub(r"[^A-Za-z0-9]", "", str(_clean_cell(asin))).upper()
    suffix = clean_asin or f"{row_number:04d}"
    return f"{niche}-{suffix}"[:50]


def prepare_products(
    source: pd.DataFrame,
    *,
    multiplier: float = 2.0,
    fixed_cost: float = 12.0,
    divisor: float = 0.8,
) -> pd.DataFrame:
    if divisor <= 0:
        raise ValueError("Hệ số chia phải lớn hơn 0.")
    source = source.copy().reset_index(drop=True)
    keyword = _source_series(source, ("keyword", "niche", "ngách"))
    asin = _source_series(source, ("asin", "amazon asin"))
    original_price = pd.to_numeric(
        _source_series(source, ("original_price", "price", "giá")),
        errors="coerce",
    )
    calculated_price = original_price.apply(
        lambda value: calculate_tiktok_price(value, multiplier, fixed_cost, divisor)
    )
    imported_final_price = pd.to_numeric(
        _source_series(source, ("tiktok_price", "final price"), default=None),
        errors="coerce",
    )
    final_price = imported_final_price.where(
        imported_final_price.notna(), calculated_price
    )

    result = pd.DataFrame(index=source.index)
    result["selected"] = (
        _source_series(source, ("selected",), True).fillna(True).astype(bool)
    )
    result["keyword"] = keyword.fillna("").astype(str)
    result["asin"] = asin.fillna("").astype(str)
    imported_sku = _source_series(source, ("seller_sku", "seller sku", "sku"))
    result["seller_sku"] = [
        str(_clean_cell(value)).strip()
        or _build_sku(result.at[index, "keyword"], result.at[index, "asin"], index + 1)
        for index, value in imported_sku.items()
    ]
    result["product_title"] = (
        _source_series(source, ("product_title", "title", "product name"))
        .fillna("")
        .astype(str)
        .str.strip()
    )
    result["description"] = (
        _source_series(source, ("description", "product description"))
        .fillna("")
        .astype(str)
        .str.strip()
    )
    result["brand"] = (
        _source_series(source, ("brand", "brand name"))
        .fillna("")
        .astype(str)
        .str.strip()
    )
    result["original_price"] = original_price
    result["calculated_price"] = calculated_price
    result["tiktok_price"] = final_price
    result["stock"] = (
        pd.to_numeric(
            _source_series(source, ("stock", "quantity", "inventory"), 1),
            errors="coerce",
        )
        .fillna(1)
        .astype(int)
    )
    result["image_1"] = (
        _source_series(source, ("image_1", "image_url", "main image"))
        .fillna("")
        .astype(str)
    )
    result["image_2"] = (
        _source_series(source, ("image_2", "product image 2"))
        .fillna("")
        .astype(str)
    )
    result["product_url"] = (
        _source_series(source, ("product_url", "amazon url", "url"))
        .fillna("")
        .astype(str)
    )
    for column, aliases in (
        ("package_weight_lb", ("package_weight_lb", "package weight", "weight")),
        ("package_length_in", ("package_length_in", "package length", "length")),
        ("package_width_in", ("package_width_in", "package width", "width")),
        ("package_height_in", ("package_height_in", "package height", "height")),
    ):
        result[column] = pd.to_numeric(
            _source_series(source, aliases, default=None), errors="coerce"
        )
    result["gtin_upc"] = (
        _source_series(source, ("gtin_upc", "gtin", "upc")).fillna("").astype(str)
    )
    result["variation_name"] = (
        _source_series(source, ("variation_name", "variation 1 name"))
        .fillna("")
        .astype(str)
    )
    result["variation_value"] = (
        _source_series(source, ("variation_value", "variation 1 value"))
        .fillna("")
        .astype(str)
    )
    result["validation"] = ""
    return validate_products(result)[EDITOR_COLUMNS]


def validate_products(frame: pd.DataFrame) -> pd.DataFrame:
    validated = frame.copy()
    messages: list[str] = []
    for _, row in validated.iterrows():
        issues: list[str] = []
        selected_value = row.get("selected", False)
        is_selected = (
            False
            if selected_value is None or pd.isna(selected_value)
            else bool(selected_value)
        )
        if not is_selected:
            messages.append("Không xuất")
            continue
        if not str(_clean_cell(row.get("seller_sku", ""))).strip():
            issues.append("Thiếu Seller SKU")
        title = str(_clean_cell(row.get("product_title", ""))).strip()
        if not title:
            issues.append("Thiếu tên")
        elif len(title) < 25:
            issues.append("Tên dưới 25 ký tự")
        elif len(title) > 200:
            issues.append("Tên trên 200 ký tự")
        if not str(_clean_cell(row.get("description", ""))).strip():
            issues.append("Thiếu mô tả")
        if not str(_clean_cell(row.get("brand", ""))).strip():
            issues.append("Chưa xác nhận brand")
        if not str(_clean_cell(row.get("image_1", ""))).strip():
            issues.append("Thiếu ảnh chính")
        price = pd.to_numeric(
            pd.Series([row.get("tiktok_price")]), errors="coerce"
        ).iloc[0]
        if pd.isna(price) or float(price) <= 0:
            issues.append("Giá không hợp lệ")
        stock = pd.to_numeric(pd.Series([row.get("stock")]), errors="coerce").iloc[0]
        if (
            pd.isna(stock)
            or float(stock) < 0
            or not float(stock).is_integer()
        ):
            issues.append("Tồn kho không hợp lệ")
        if any(
            pd.isna(
                pd.to_numeric(pd.Series([row.get(column)]), errors="coerce").iloc[0]
            )
            for column in (
                "package_weight_lb",
                "package_length_in",
                "package_width_in",
                "package_height_in",
            )
        ):
            issues.append("Thiếu cân nặng/kích thước")
        messages.append("Sẵn sàng" if not issues else "; ".join(issues))
    validated["validation"] = messages
    selected_mask = validated["selected"].fillna(False).astype(bool)
    normalized_skus = (
        validated["seller_sku"].fillna("").astype(str).str.strip().str.casefold()
    )
    duplicate_mask = selected_mask & normalized_skus.ne("") & normalized_skus.duplicated(
        keep=False
    )
    for index in validated.index[duplicate_mask]:
        current = str(validated.at[index, "validation"])
        validated.at[index, "validation"] = (
            "Trùng Seller SKU"
            if current == "Sẵn sàng"
            else f"{current}; Trùng Seller SKU"
        )
    return validated


def selected_products(frame: pd.DataFrame) -> pd.DataFrame:
    if "selected" not in frame.columns:
        return frame.copy()
    return frame.loc[frame["selected"].fillna(False).astype(bool)].reset_index(drop=True)


def preparation_filename(frame: pd.DataFrame) -> str:
    if not frame.empty and "keyword" in frame.columns:
        first_keyword = str(frame.iloc[0].get("keyword", "")).strip()
    else:
        first_keyword = "tiktok_products"
    return f"{slugify_filename(first_keyword or 'tiktok_products')}_tiktok_draft.xlsx"


def build_preparation_workbook(
    frame: pd.DataFrame,
    *,
    multiplier: float = 2.0,
    fixed_cost: float = 12.0,
    divisor: float = 0.8,
) -> bytes:
    products = validate_products(selected_products(frame))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TikTok Products"
    settings = workbook.create_sheet("Pricing Settings")

    settings.append(["Thiết lập", "Giá trị", "Ghi chú"])
    settings.append(["Hệ số nhân giá gốc", float(multiplier), "Mặc định theo yêu cầu"])
    settings.append(["Chi phí cộng thêm", float(fixed_cost), "USD"])
    settings.append(["Hệ số chia", float(divisor), "Phải lớn hơn 0"])
    settings.append(
        [
            "Công thức",
            "Giá bán = (Giá gốc × Hệ số nhân + Chi phí cộng thêm) ÷ Hệ số chia",
            "Giá bán cuối có thể sửa thủ công",
        ]
    )

    headers = [
        "Xuất",
        "Ngách",
        "ASIN",
        "Seller SKU",
        "Tên sản phẩm",
        "Mô tả",
        "Brand",
        "Giá gốc",
        "Giá công thức",
        "Giá bán TikTok",
        "Tồn kho",
        "Ảnh chính",
        "Ảnh 2",
        "Link nguồn",
        "Cân nặng (lb)",
        "Dài (in)",
        "Rộng (in)",
        "Cao (in)",
        "GTIN/UPC",
        "Tên biến thể 1",
        "Giá trị biến thể 1",
        "Kiểm tra",
    ]
    sheet.append(headers)
    for row_index, (_, product) in enumerate(products.iterrows(), start=2):
        values = [
            bool(product.get("selected", True)),
            _clean_cell(product.get("keyword")),
            _clean_cell(product.get("asin")),
            _clean_cell(product.get("seller_sku")),
            _clean_cell(product.get("product_title")),
            _clean_cell(product.get("description")),
            _clean_cell(product.get("brand")),
            _clean_cell(product.get("original_price")),
            None,
            _clean_cell(product.get("tiktok_price")),
            _clean_cell(product.get("stock")),
            _clean_cell(product.get("image_1")),
            _clean_cell(product.get("image_2")),
            _clean_cell(product.get("product_url")),
            _clean_cell(product.get("package_weight_lb")),
            _clean_cell(product.get("package_length_in")),
            _clean_cell(product.get("package_width_in")),
            _clean_cell(product.get("package_height_in")),
            _clean_cell(product.get("gtin_upc")),
            _clean_cell(product.get("variation_name")),
            _clean_cell(product.get("variation_value")),
            _clean_cell(product.get("validation")),
        ]
        sheet.append(values)
        sheet.cell(row_index, 9).value = (
            f"=ROUND((H{row_index}*'Pricing Settings'!$B$2+"
            f"'Pricing Settings'!$B$3)/'Pricing Settings'!$B$4,2)"
        )

    orange = "D97706"
    navy = "172554"
    cream = "FFF7ED"
    green = "DCFCE7"
    warning = "FEF3C7"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="E5E7EB")
    for target in (sheet, settings):
        target.sheet_view.showGridLines = False
        target.freeze_panes = "A2"
        for cell in target[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="medium", color=orange))
        target.row_dimensions[1].height = 28

    for row in sheet.iter_rows(min_row=2, max_row=max(sheet.max_row, 2)):
        for cell in row:
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(
                vertical="top", wrap_text=cell.column in {5, 6, 22}
            )
        status = str(sheet.cell(cell.row, 22).value or "")
        sheet.cell(cell.row, 22).fill = PatternFill(
            "solid", fgColor=green if status == "Sẵn sàng" else warning
        )
    for column in (8, 9, 10):
        for cell in sheet.iter_cols(
            min_col=column,
            max_col=column,
            min_row=2,
            max_row=max(sheet.max_row, 2),
        ):
            for item in cell:
                item.number_format = '$#,##0.00'
    for column in range(1, settings.max_column + 1):
        settings.cell(1, column).fill = PatternFill("solid", fgColor=orange)
    settings["B2"].fill = PatternFill("solid", fgColor=cream)
    settings["B3"].fill = PatternFill("solid", fgColor=cream)
    settings["B4"].fill = PatternFill("solid", fgColor=cream)

    widths = {
        1: 9,
        2: 18,
        3: 14,
        4: 24,
        5: 44,
        6: 52,
        7: 20,
        8: 14,
        9: 16,
        10: 16,
        11: 11,
        12: 42,
        13: 42,
        14: 42,
        15: 15,
        16: 12,
        17: 12,
        18: 12,
        19: 18,
        20: 20,
        21: 22,
        22: 42,
    }
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    settings.column_dimensions["A"].width = 28
    settings.column_dimensions["B"].width = 62
    settings.column_dimensions["C"].width = 34
    sheet.auto_filter.ref = sheet.dimensions

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def workbook_sheet_names(template_bytes: bytes) -> list[str]:
    workbook = load_workbook(
        io.BytesIO(template_bytes), read_only=True, data_only=False
    )
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def detect_header_row(
    template_bytes: bytes, sheet_name: str, max_scan_rows: int = 30
) -> int:
    workbook = load_workbook(
        io.BytesIO(template_bytes), read_only=True, data_only=False
    )
    try:
        sheet = workbook[sheet_name]
        aliases = {
            _normalized_name(alias)
            for field in EXPORT_FIELDS
            for alias in field.aliases
        }
        best_row = 1
        best_score = -1
        for row_number in range(1, min(sheet.max_row, max_scan_rows) + 1):
            values = [
                _normalized_name(cell.value)
                for cell in sheet[row_number]
                if cell.value not in (None, "")
            ]
            alias_hits = sum(value in aliases for value in values)
            score = alias_hits * 20 + len(values)
            if score > best_score:
                best_row = row_number
                best_score = score
        return best_row
    finally:
        workbook.close()


def template_headers(
    template_bytes: bytes,
    sheet_name: str,
    header_row: int,
) -> list[dict[str, Any]]:
    workbook = load_workbook(
        io.BytesIO(template_bytes), read_only=True, data_only=False
    )
    try:
        sheet = workbook[sheet_name]
        headers: list[dict[str, Any]] = []
        seen: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            value = str(sheet.cell(header_row, column).value or "").strip()
            if not value:
                continue
            seen[value] = seen.get(value, 0) + 1
            label = (
                value
                if seen[value] == 1
                else f"{value} [{get_column_letter(column)}]"
            )
            headers.append({"label": label, "header": value, "column": column})
        return headers
    finally:
        workbook.close()


def suggest_template_mapping(
    headers: list[Mapping[str, Any]],
) -> dict[str, str]:
    result: dict[str, str] = {}
    normalized_headers = {
        _normalized_name(item.get("header", "")): str(item.get("label", ""))
        for item in headers
    }
    for field in EXPORT_FIELDS:
        result[field.key] = next(
            (
                normalized_headers[_normalized_name(alias)]
                for alias in field.aliases
                if _normalized_name(alias) in normalized_headers
            ),
            "",
        )
    return result


def fill_official_template(
    template_bytes: bytes,
    *,
    sheet_name: str,
    header_row: int,
    mapping: Mapping[str, str],
    products: pd.DataFrame,
) -> bytes:
    workbook = load_workbook(io.BytesIO(template_bytes), data_only=False)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError("Không tìm thấy sheet đã chọn trong template.")
    sheet = workbook[sheet_name]
    headers = template_headers(template_bytes, sheet_name, header_row)
    columns_by_label = {
        str(item["label"]): int(item["column"]) for item in headers
    }
    selected = selected_products(products)
    if selected.empty:
        workbook.close()
        raise ValueError("Chưa chọn sản phẩm nào để xuất.")

    used_targets: set[int] = set()
    clean_mapping: dict[str, int] = {}
    for source_key, target_label in mapping.items():
        if not target_label:
            continue
        target_column = columns_by_label.get(str(target_label))
        if target_column is None:
            workbook.close()
            raise ValueError(f"Cột template không còn tồn tại: {target_label}")
        if target_column in used_targets:
            workbook.close()
            raise ValueError("Mỗi cột template chỉ được ánh xạ một lần.")
        used_targets.add(target_column)
        clean_mapping[source_key] = target_column
    if not clean_mapping:
        workbook.close()
        raise ValueError("Chưa ánh xạ cột nào vào template TikTok.")

    style_row = header_row + 1
    original_max_row = sheet.max_row
    for offset, (_, product) in enumerate(selected.iterrows(), start=1):
        target_row = header_row + offset
        for source_key, target_column in clean_mapping.items():
            target_cell = sheet.cell(target_row, target_column)
            if target_row > original_max_row or not target_cell.has_style:
                style_cell = sheet.cell(style_row, target_column)
                if style_cell.has_style:
                    target_cell._style = copy(style_cell._style)
                    target_cell.number_format = style_cell.number_format
                    target_cell.alignment = copy(style_cell.alignment)
                    target_cell.protection = copy(style_cell.protection)
            value = _clean_cell(product.get(source_key, ""))
            if isinstance(value, float) and pd.isna(value):
                value = ""
            target_cell.value = value

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
